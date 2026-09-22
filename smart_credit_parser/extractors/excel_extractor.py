"""
Excel extractor supporting .xlsx files via openpyxl.
"""

from __future__ import annotations
from pathlib import Path
from typing import Any, Dict, Iterator, List, Tuple
import openpyxl
from .base import BaseExtractor


class ExcelExtractor(BaseExtractor):
    """Extractor for Microsoft Excel (.xlsx, .xlsm) files."""

    SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}

    @classmethod
    def can_handle(cls, file_path: Path) -> bool:
        return file_path.suffix.lower() in cls.SUPPORTED_EXTENSIONS

    def get_sample(
        self, file_path: Path, max_rows: int = 20
    ) -> Tuple[List[str], List[Dict[str, Any]]]:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheet = wb.active
        if not sheet:
            wb.close()
            return [], []

        rows_iter = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows_iter)
        except StopIteration:
            wb.close()
            return [], []

        headers = [
            str(h).strip() if h is not None else f"col_{idx}"
            for idx, h in enumerate(raw_headers)
        ]

        rows: List[Dict[str, Any]] = []
        for row in rows_iter:
            if not row or not any(x is not None for x in row):
                continue
            row_dict = {}
            for idx, h in enumerate(headers):
                val = row[idx] if idx < len(row) else None
                row_dict[h] = "" if val is None else str(val).strip()
            rows.append(row_dict)
            if len(rows) >= max_rows:
                break

        wb.close()
        return headers, rows

    def extract_chunks(
        self, file_path: Path, chunk_size: int = 1000
    ) -> Iterator[List[Dict[str, Any]]]:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheet = wb.active
        if not sheet:
            wb.close()
            return

        rows_iter = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows_iter)
        except StopIteration:
            wb.close()
            return

        headers = [
            str(h).strip() if h is not None else f"col_{idx}"
            for idx, h in enumerate(raw_headers)
        ]

        chunk: List[Dict[str, Any]] = []
        for row in rows_iter:
            if not row or not any(x is not None for x in row):
                continue
            row_dict = {}
            for idx, h in enumerate(headers):
                val = row[idx] if idx < len(row) else None
                row_dict[h] = "" if val is None else str(val).strip()
            chunk.append(row_dict)
            if len(chunk) >= chunk_size:
                yield chunk
                chunk = []

        if chunk:
            yield chunk

        wb.close()
